"""Synthetic acceptance operations using each repository's own installed runtime."""

from pathlib import Path
import csv
import hashlib
import secrets
import io
import json
import os
import shutil
import sqlite3
import subprocess
import sys
import time
import traceback
import zipfile

ROOT = Path(__file__).resolve().parents[1]
os.chdir(ROOT)
sys.path.insert(0, str(ROOT))
from app.service import Application
from localdesk.safety import InputError
from tests.fixtures import make_image, make_pdf

import argparse

parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument(
    "--report-dir",
    type=Path,
    required=True,
    help="New ignored evidence folder for this run. Existing folders are preserved.",
)
args = parser.parse_args()
BASE = args.report_dir.resolve()
BASE.mkdir(parents=True, exist_ok=False)
WORK = BASE / "synthetic space unicode-ÃƒÂ©"
WORK.mkdir()
RESULTS = []
APP = None
PASSWORD = secrets.token_urlsafe(32)


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def record(name, func):
    started = time.monotonic()
    try:
        detail = func()
        result = dict(check=name, status="PASS", detail=detail)
    except Exception:
        result = dict(check=name, status="FAIL", traceback=traceback.format_exc())
    result["seconds"] = time.monotonic() - started
    RESULTS.append(result)
    (BASE / "results.json").write_text(
        json.dumps(RESULTS, indent=2, default=str), encoding="utf-8"
    )
    print(name, result["status"], flush=True)


def finish(action, body, expected="done"):
    ident = APP.common_post(action, body)["job_id"]
    deadline = time.monotonic() + 180
    while time.monotonic() < deadline:
        job = APP.jobs.get(ident)
        if job["status"] not in ("running", "queued"):
            assert job["status"] == expected, job
            return job["result"] if expected == "done" else job
        time.sleep(0.03)
    raise TimeoutError(action)


def output(artifact):
    return APP.download(artifact["path"]).read_bytes()


def restart():
    global APP
    APP.close()
    APP = Application(
        ROOT,
        BASE / "persistent-data",
        passphrase=PASSWORD,
    )


def file(name, content):
    path = WORK / name
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(content.encode("utf-8") if isinstance(content, str) else content)
    return path


def office_fixtures():
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation
    from pptx.util import Inches

    doc = Document()
    doc.add_paragraph("Office canary qa@example.test public number 42")
    doc.core_properties.author = "PRIVATE_AUTHOR_CANARY"
    doc.save(WORK / "office.docx")
    book = Workbook()
    book.active.append(["Office canary qa@example.test", 42])
    book.save(WORK / "office.xlsx")
    book.close()
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text = (
        "Office canary qa@example.test public number 42"
    )
    deck.save(WORK / "office.pptx")
    return [WORK / ("office" + ext) for ext in (".docx", ".xlsx", ".pptx")]


def timeline():
    def events():
        p = file("history.txt", "PRIVATE_DOCUMENT_CONTENT_CANARY")
        APP.common_post("roots/add", {"path": str(WORK), "project": "QA"})
        finish("scan", {})
        assert APP.get("events", {})["events"][0]["kind"] == "discovered"
        created = file("new.txt", "PRIVATE_DOCUMENT_CONTENT_CANARY")
        finish("scan", {})
        created.write_text("Different length PRIVATE_DOCUMENT_CONTENT_CANARY")
        finish("scan", {})
        created.rename(WORK / "renamed.txt")
        finish("scan", {})
        (WORK / "renamed.txt").unlink()
        finish("scan", {})
        rows = APP.get("events", {})["events"]
        assert {"discovered", "created", "modified", "deleted"} <= {
            r["kind"] for r in rows
        }
        assert "PRIVATE_DOCUMENT_CONTENT_CANARY" not in json.dumps(rows)
        APP.common_post(
            "event/update",
            {"id": rows[0]["id"], "note": "Reviewed output", "tags": ["acceptance"]},
        )
        APP.common_post("note", {"note": "Synthetic note", "project": "QA"})
        graph = APP.get("graph", {})
        return {"events": rows, "graph": graph}

    record(
        "AG-01/AG-02/AG-03 discovery lifecycle content exclusion notes tags graph",
        events,
    )

    def persist():
        APP.common_post("monitor/start", {"seconds": 15})
        restart()
        assert APP.state()["monitor"]["enabled"]
        assert not APP.state()["capture"]["enabled"]
        file("background.txt", "PRIVATE_DOCUMENT_CONTENT_CANARY")
        deadline = time.monotonic() + 40
        while time.monotonic() < deadline:
            if any(
                "background.txt" in r["path"] for r in APP.get("events", {})["events"]
            ):
                break
            time.sleep(0.3)
        else:
            raise AssertionError("Restarted monitor missed new file")
        APP.common_post("monitor/stop", {})
        restart()
        assert not APP.state()["monitor"]["enabled"]
        return APP.state()

    record(
        "AG-07 persistent monitor and paused restart real watched-file event", persist
    )


def shared():
    def persist():
        APP.store.set("qa_canary", "PERSISTED_SYNTHETIC_CANARY")
        backup = APP.common_post("backup", {})
        raw = output(backup)
        assert b"PERSISTED_SYNTHETIC_CANARY" not in raw
        restart()
        assert APP.store.get("qa_canary") == "PERSISTED_SYNTHETIC_CANARY"
        from localdesk.storage import Store

        restored = BASE / "restored.vault"
        restored.write_bytes(raw)
        store = Store(restored, PASSWORD)
        assert store.get("qa_canary") == "PERSISTED_SYNTHETIC_CANARY"
        store.close()
        before = sha(restored)
        try:
            Store(restored, secrets.token_urlsafe(32))
        except InputError:
            pass
        else:
            raise AssertionError("Wrong password accepted")
        assert sha(restored) == before
        return {
            "backup_sha256": hashlib.sha256(raw).hexdigest(),
            "restart_and_restore": True,
        }

    record(
        "Shared real encrypted persistence backup restore wrong-password preservation",
        persist,
    )


def main():
    global APP
    APP = Application(
        ROOT,
        BASE / "persistent-data",
        passphrase=PASSWORD,
    )
    try:
        timeline()
        shared()
    finally:
        APP.close()
        hashes = {
            p.relative_to(BASE).as_posix(): sha(p)
            for p in BASE.rglob("*")
            if p.is_file() and p.name != "hashes.json"
        }
        (BASE / "hashes.json").write_text(
            json.dumps(hashes, indent=2), encoding="utf-8"
        )
    return int(any(r["status"] == "FAIL" for r in RESULTS))


if __name__ == "__main__":
    raise SystemExit(main())
